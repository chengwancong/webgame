# 配置表导出：config/excel/*.xlsx → config/json/*.json
# 运行：pnpm config:export；pnpm config:check 只比对，不写文件（不一致时退出码 1）。
# 表头约定见每本工作簿的「说明」页签：第1行中文名，第2行字段键，第3行类型，第4行起数据。
import json, os, sys
from openpyxl import load_workbook

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'config')
EXCEL, JSON_DIR = os.path.join(ROOT, 'excel'), os.path.join(ROOT, 'json')
TYPES = {'string', 'int', 'number', 'bool', 'list', 'json'}
TRUE, FALSE = {'是', 'true', 'TRUE', 'True', '1', 'y', 'yes'}, {'否', 'false', 'FALSE', 'False', '0', 'n', 'no'}

class ConfigError(Exception): pass

def parse(t, v, where):
  if isinstance(v, str): v = v.strip()
  if v is None or v == '': return None
  try:
    if t == 'string':
      if isinstance(v, float) and v.is_integer(): v = int(v)
      return str(v)
    if t == 'int':
      if isinstance(v, bool): raise ValueError
      f = float(v)
      if not f.is_integer(): raise ValueError
      return int(f)
    if t == 'number':
      if isinstance(v, bool): raise ValueError
      f = float(v)
      return int(f) if f.is_integer() else f
    if t == 'bool':
      if isinstance(v, bool): return v
      s = str(v)
      if s in TRUE: return True
      if s in FALSE: return False
      raise ValueError
    if t == 'list': return [x.strip() for x in str(v).split('|') if x.strip()]
    if t == 'json': return json.loads(v) if isinstance(v, str) else v
  except (ValueError, TypeError, json.JSONDecodeError):
    raise ConfigError(f'{where}：值 {v!r} 不是有效的 {t}')
  raise ConfigError(f'{where}：未知类型 {t}')

def read_sheet(ws, book, header=1):
  keys = [c.value for c in ws[header+1]]; types = [c.value for c in ws[header+2]]; names = [c.value for c in ws[header]]
  cols = []
  for i, k in enumerate(keys):
    if k is None or str(k).strip() == '': continue
    k = str(k).strip(); t = str(types[i] or '').strip()
    if t not in TYPES: raise ConfigError(f'{book}/{ws.title} 第3行 {names[i] or k}：未知类型 {t!r}')
    if any(c[1] == k for c in cols): raise ConfigError(f'{book}/{ws.title}：字段键 {k} 重复')
    cols.append((i, k, t, names[i] or k))
  rows = []
  for r, row in enumerate(ws.iter_rows(min_row=header+3, values_only=True), header+3):
    first = row[0] if row else None
    if isinstance(first, str) and first.strip().startswith('#'): continue
    out = {}
    for i, k, t, n in cols:
      v = parse(t, row[i] if i < len(row) else None, f'{book}/{ws.title} 第{r}行「{n}」')
      if v is not None: out[k] = v
    if out: out['_row'] = r; rows.append(out)
  return rows

def export_book(path):
  book = os.path.basename(path)
  wb = load_workbook(path, data_only=True)
  if '导出' not in wb.sheetnames: raise ConfigError(f'{book}：缺少「导出」页签')
  ex = wb['导出']; output = str(ex['B1'].value or '').strip()
  if not output.endswith('.json'): raise ConfigError(f'{book}/导出 B1：输出文件须为 .json')
  data = {'source': 'config/excel/' + book}
  for r in ex.iter_rows(min_row=3, values_only=True):
    if not r or not r[0]: continue
    title, key = str(r[0]).strip(), str(r[1] or '').strip()
    if title not in wb.sheetnames: raise ConfigError(f'{book}/导出：页签 {title} 不存在')
    if not key or key in data: raise ConfigError(f'{book}/导出：页签 {title} 的导出键为空或重复')
    data[key] = read_sheet(wb[title], book)
  if book == '任务.xlsx':
    guide = wb['说明']
    marker = next((r for r in range(1, guide.max_row+1) if guide.cell(r,1).value == '#quest-parameters'), None)
    if marker: data['parameters'] = read_sheet(guide, book, marker+1)
  wb.close()
  return output, data

def render(data):
  parts = ['{', f' "source": {json.dumps(data["source"], ensure_ascii=False)}']
  for k, rows in data.items():
    if k == 'source': continue
    body = ',\n'.join('  ' + json.dumps(row, ensure_ascii=False, separators=(',', ':')) for row in rows)
    parts[-1] += ','
    parts.append(f' {json.dumps(k)}: [\n{body}\n ]' if rows else f' {json.dumps(k)}: []')
  parts.append('}')
  return '\n'.join(parts) + '\n'

def main():
  check = '--check' in sys.argv
  books = sorted(f for f in os.listdir(EXCEL) if f.endswith('.xlsx') and not f.startswith('~$'))
  if not books: raise ConfigError('config/excel 下没有工作簿')
  os.makedirs(JSON_DIR, exist_ok=True)
  stale, seen, exported = [], {}, {}
  for f in books:
    output, data = export_book(os.path.join(EXCEL, f))
    if output in seen: raise ConfigError(f'{f} 与 {seen[output]} 输出到同一个 {output}')
    seen[output] = f
    exported[output] = data
    if f == '制作.xlsx':
      for name, selected in [('crafting-public.json', {k: v for k, v in data.items() if k != 'drops'}), ('crafting-private.json', {'source': data['source'], 'drops': data['drops']})]:
        seen[name] = f
        exported[name] = selected
    # 机制旋钮独立导出，保持用户的 builds/weights/settings 原样。
    if f == '节奏目标.xlsx':
      wb = load_workbook(os.path.join(EXCEL, f), data_only=True)
      if '机制旋钮' in wb.sheetnames:
        knob_output = 'pacing-mechanism-knobs.json'
        seen[knob_output] = f
        exported[knob_output] = {'source': 'config/excel/' + f, 'knobs': read_sheet(wb['机制旋钮'], f)}
      if '主干目标' in wb.sheetnames:
        spine_output = 'spine-targets.json'
        seen[spine_output] = f
        exported[spine_output] = {'source': 'config/excel/' + f, 'targets': read_sheet(wb['主干目标'], f)}
      wb.close()
  from config_quest_check import validate_quest_config
  validate_quest_config(exported)
  from config_crafting_check import validate_crafting_config
  try: validate_crafting_config(exported)
  except (ValueError, KeyError, TypeError) as error: raise ConfigError(str(error))
  for output, data in exported.items():
    f = seen[output]
    text, target = render(data), os.path.join(JSON_DIR, output)
    old = open(target, encoding='utf-8').read() if os.path.exists(target) else None
    if old == text: continue
    if check: stale.append(output)
    else:
      with open(target, 'w', encoding='utf-8', newline='\n') as fh: fh.write(text)
      print('导出', f, '→', output)
  if check:
    if stale: print('以下 JSON 与 Excel 不一致，请运行 pnpm config:export：', ', '.join(stale)); sys.exit(1)
    print('配置表一致：', len(books), '本工作簿')

if __name__ == '__main__':
  try: main()
  except ConfigError as e: print('配置表错误：', e, file=sys.stderr); sys.exit(1)
