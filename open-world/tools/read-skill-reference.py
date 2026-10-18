"""Read-only extraction of the user's combat workbooks, retaining cell addresses."""
from pathlib import Path
import hashlib,json
import openpyxl

root=Path('G:/webgame/abyss-walker-migration/excel')
names=['战斗_技能.xlsx','战斗_BUFF.xlsx','战斗_常量.xlsx','战斗_属性.xlsx','角色_技能.xlsx','角色_技能BUFF.xlsx','新增技能.xlsx','辅助技能配置设计.xlsx','副本_BUFF机制.xlsx','技能表-text.xlsx','战斗技能对比用.xlsx']
result=[]
for name in names:
    path=root/name
    if not path.exists():continue
    book=openpyxl.load_workbook(path,read_only=True,data_only=False)
    record={'path':str(path),'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'sheets':[]}
    for sheet in book:
        rows=[]
        for index,row in enumerate(sheet.iter_rows(),1):
            values={cell.coordinate:cell.value for cell in row if cell.value is not None}
            if values:rows.append({'row':index,'cells':values})
        record['sheets'].append({'name':sheet.title,'rows':rows,'maxRow':sheet.max_row,'maxColumn':sheet.max_column})
    result.append(record)
    print(json.dumps({'file':name,'sheets':[{'name':s['name'],'rows':s['maxRow'],'cols':s['maxColumn'],'head':s['rows'][:3]} for s in record['sheets']]},ensure_ascii=False,default=str))
    book.close()
destination=Path('docs/skill-repair-20261003/excel-reference.json')
destination.write_text(json.dumps(result,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
